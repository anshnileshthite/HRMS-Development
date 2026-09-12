import openpyxl
import pandas as pd
import hashlib
import os
from sqlalchemy import create_engine, text
from sqlalchemy.engine import URL

SUPABASE_URL = URL.create(
    drivername="postgresql+pg8000",
    username="postgres.lsyfgompyfjborgxzxkx",
    password="Haveaniceday@2027",
    host="aws-0-ap-southeast-1.pooler.supabase.com",
    port=6543,
    database="postgres"
)

# Define each firm and its corresponding wage sheet
FIRMS_CONFIG = [
    {
        "company_id": "SAGAR",
        "company_name": "SAGAR ENTERPRISES",
        "file_name": "wages_aug_2026.xlsx",
        "start_row": 6
    },
    {
        "company_id": "GEMSHINE",
        "company_name": "GEMSHINE MULTISERVICES",
        "file_name": "Gemshine contract Aug  2026 wages sheet.xlsx",
        "start_row": 4
    },
    {
        "company_id": "ELITE",
        "company_name": "ELITE MULTISERVICES",
        "file_name": "Elite contract Aug 2026 wages sheet.xlsx", # Place when ready
        "start_row": 5
    }
]

def parse_and_sync():
    engine = create_engine(SUPABASE_URL)

    with engine.connect() as conn:
        for firm in FIRMS_CONFIG:
            f_id = firm["company_id"]
            f_name = firm["file_name"]
            start_row = firm["start_row"]

            if not os.path.exists(f_name):
                print(f"Skipping {firm['company_name']} ({f_name} not found in folder).")
                continue

            print(f"\nProcessing {firm['company_name']} from '{f_name}'...")
            wb = openpyxl.load_workbook(f_name, data_only=True)
            ws = wb.active

            row = start_row
            count = 0
            while row < 5000:
                sr_val = ws.cell(row=row, column=1).value
                emp_id_val = ws.cell(row=row, column=2).value

                if not emp_id_val or "total" in str(sr_val or "").lower():
                    if ws.cell(row=row + 1, column=2).value is None:
                        break
                    row += 1
                    continue

                emp_code = str(emp_id_val).strip().replace(".0", "")
                if not emp_code:
                    row += 1
                    continue

                def c_str(c): return str(ws.cell(row=row, column=c).value or "").strip()
                def c_num(c):
                    v = ws.cell(row=row, column=c).value
                    try: return float(str(v).replace(',', '').strip()) if v is not None else 0.0
                    except: return 0.0

                # 1. Sync Employee
                conn.execute(text("""
                    INSERT INTO employees (company_id, emp_code, full_name, father_name, aadhaar_no, uan_no)
                    VALUES (:cid, :code, :name, :fname, :adh, :uan)
                    ON CONFLICT (company_id, emp_code) DO UPDATE SET
                        full_name = EXCLUDED.full_name,
                        father_name = EXCLUDED.father_name;
                """), {
                    "cid": f_id,
                    "code": emp_code,
                    "name": c_str(3),
                    "fname": c_str(4),
                    "adh": c_str(5),
                    "uan": c_str(6)
                })

                # 2. Sync Credentials (EmpCode@123)
                pwd_hash = hashlib.sha256(f"{emp_code}@123".encode()).hexdigest()
                conn.execute(text("""
                    INSERT INTO user_accounts (company_id, emp_code, password_hash, role)
                    VALUES (:cid, :code, :pwd, 'employee')
                    ON CONFLICT (company_id, emp_code) DO NOTHING;
                """), {"cid": f_id, "code": emp_code, "pwd": pwd_hash})

                count += 1
                row += 1

            # Create Super Admin login for each firm
            admin_pwd = hashlib.sha256(f"{f_id}@2026".encode()).hexdigest()
            conn.execute(text("""
                INSERT INTO user_accounts (company_id, emp_code, password_hash, role)
                VALUES (:cid, 'ADMIN', :pwd, 'admin')
                ON CONFLICT (company_id, emp_code) DO UPDATE SET password_hash = EXCLUDED.password_hash;
            """), {"cid": f_id, "pwd": admin_pwd})

            conn.commit()
            print(f"Synced {count} workers for {firm['company_name']}!")

if __name__ == '__main__':
    parse_and_sync()