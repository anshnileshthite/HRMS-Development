import sqlite3
import openpyxl
import hashlib
import os

def init_and_populate_db():
    conn = sqlite3.connect('khed_payroll.db')
    cursor = conn.cursor()

    # 1. Employees Master
    cursor.execute('''
    CREATE TABLE IF NOT EXISTS employees (
        emp_code TEXT PRIMARY KEY,
        full_name TEXT NOT NULL,
        father_name TEXT,
        uan_no TEXT,
        esic_no TEXT,
        pan_no TEXT,
        aadhaar_no TEXT,
        gender TEXT,
        bank_name TEXT,
        bank_account TEXT,
        bank_ifsc TEXT,
        mobile TEXT,
        category TEXT
    )''')

    # 2. Monthly Processed Wages
    cursor.execute('''
    CREATE TABLE IF NOT EXISTS monthly_wages (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        emp_code TEXT,
        month INTEGER,
        year INTEGER,
        rate_basic REAL,
        rate_da REAL,
        rate_hra REAL,
        rate_wages REAL,
        rate_ot REAL,
        days_worked REAL,
        paid_holidays REAL,
        total_days REAL,
        ot_hours REAL,
        earned_basic REAL,
        earned_da REAL,
        earned_hra REAL,
        earned_wages REAL,
        earned_ot_amt REAL,
        gross_amount REAL,
        ded_pf REAL,
        ded_esic REAL,
        ded_pt REAL,
        ded_mlwf REAL,
        ded_advance REAL,
        total_deduction REAL,
        net_wages REAL,
        er_pf REAL,
        er_esic REAL,
        service_charge REAL,
        sub_total REAL,
        gst_18 REAL,
        ctc_total REAL,
        UNIQUE (emp_code, month, year)
    )''')

    # 3. User Accounts
    cursor.execute('''
    CREATE TABLE IF NOT EXISTS user_accounts (
        emp_code TEXT PRIMARY KEY,
        password_hash TEXT NOT NULL,
        role TEXT NOT NULL DEFAULT 'employee'
    )''')

    # 4. Daily Punches
    cursor.execute('''
    CREATE TABLE IF NOT EXISTS daily_punches (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        emp_code TEXT NOT NULL,
        punch_date TEXT NOT NULL,
        punch_in TEXT,
        punch_out TEXT,
        UNIQUE (emp_code, punch_date)
    )''')

    # 5. Leave Requests
    cursor.execute('''
    CREATE TABLE IF NOT EXISTS leave_requests (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        emp_code TEXT NOT NULL,
        leave_type TEXT NOT NULL,
        from_date TEXT NOT NULL,
        to_date TEXT NOT NULL,
        reason TEXT,
        status TEXT DEFAULT 'PENDING'
    )''')

    # Find the Excel wage sheet in the folder
    excel_candidates = [
        "wages_aug_2026.xlsx",
        "DAS Khed city Wages sheet Aug -2026.xlsx"
    ]
    target_excel = None
    for cand in excel_candidates:
        if os.path.exists(cand):
            target_excel = cand
            break

    if not target_excel:
        for f in os.listdir('.'):
            if f.endswith('.xlsx') and not f.startswith('~$'):
                target_excel = f
                break

    if not target_excel:
        print("No Excel file found! Please put your wage sheet (.xlsx) in this folder.")
        conn.commit()
        conn.close()
        return

    print(f"Loading data from '{target_excel}'...")
    wb = openpyxl.load_workbook(target_excel, data_only=True)
    ws = wb.active

    count = 0
    row = 5
    while row < 5000:
        sr_val = ws.cell(row=row, column=1).value
        emp_id_val = ws.cell(row=row, column=2).value

        if "total" in str(sr_val).lower() or (sr_val is None and emp_id_val is None):
            if ws.cell(row=row + 1, column=2).value is None:
                break
            row += 1
            continue

        emp_code = str(emp_id_val or "").strip()
        if not emp_code:
            row += 1
            continue

        def c_str(c): return str(ws.cell(row=row, column=c).value or "").strip()
        def c_num(c):
            v = ws.cell(row=row, column=c).value
            try: return float(str(v).replace(',', '').strip()) if v is not None else 0.0
            except: return 0.0

        # Insert Employee Master
        cursor.execute('''
            INSERT OR REPLACE INTO employees (
                emp_code, full_name, father_name, uan_no, esic_no, pan_no,
                aadhaar_no, gender, bank_name, bank_account, bank_ifsc, mobile, category
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (emp_code, c_str(3), c_str(4), c_str(5), c_str(6), c_str(7), c_str(8), c_str(9) or 'M',
              c_str(12), c_str(13), c_str(14), c_str(15), c_str(16)))

        # Insert August 2026 Wage Record
        cursor.execute('''
            INSERT OR REPLACE INTO monthly_wages (
                emp_code, month, year, rate_basic, rate_da, rate_hra, rate_wages, rate_ot,
                days_worked, paid_holidays, total_days, ot_hours,
                earned_basic, earned_da, earned_hra, earned_wages, earned_ot_amt, gross_amount,
                ded_pf, ded_esic, ded_pt, ded_mlwf, ded_advance, total_deduction, net_wages,
                er_pf, er_esic, service_charge, sub_total, gst_18, ctc_total
            ) VALUES (?, 8, 2026, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            emp_code, c_num(17), c_num(18), c_num(19), c_num(20), c_num(21),
            c_num(22), c_num(23), c_num(24), c_num(25), c_num(26), c_num(27), c_num(28), c_num(29), c_num(30), c_num(31),
            c_num(32), c_num(33), c_num(34), c_num(35), c_num(36), c_num(37), c_num(38),
            c_num(39), c_num(40), c_num(42), c_num(43), c_num(44), c_num(45)
        ))

        # Create Login Credentials (Password: EmpCode@123)
        pwd_hash = hashlib.sha256(f"{emp_code}@123".encode()).hexdigest()
        cursor.execute('''
            INSERT OR REPLACE INTO user_accounts (emp_code, password_hash, role)
            VALUES (?, ?, 'employee')
        ''', (emp_code, pwd_hash))

        count += 1
        row += 1

    # Insert Admin user
    admin_hash = hashlib.sha256("Admin@123".encode()).hexdigest()
    cursor.execute('''
        INSERT OR REPLACE INTO user_accounts (emp_code, password_hash, role)
        VALUES ('ADMIN', ?, 'admin')
    ''', (admin_hash,))

    conn.commit()
    conn.close()
    print(f"Done! Created 'khed_payroll.db' with {count} employee accounts.")

if __name__ == '__main__':
    init_and_populate_db()